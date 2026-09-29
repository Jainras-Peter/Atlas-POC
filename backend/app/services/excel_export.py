from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from app.models.quote import QuoteOut

HEADERS = [
    "Reference",
    "Type",
    "Contact Email",
    "Customer Name",
    "Mode",
    "Origin",
    "Destination",
    "Cargo",
    "Cut Off",
    "Status",
]


def quotes_to_xlsx(quotes: list[QuoteOut]) -> BytesIO:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Quotes"

    sheet.append(HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)

    for quote in quotes:
        sheet.append(
            [
                quote.quote_number,
                quote.type,
                quote.contact_email,
                quote.customer_name,
                quote.mode,
                quote.origin,
                quote.destination,
                quote.cargo,
                quote.cut_off_date,
                quote.status,
            ]
        )

    for index, header in enumerate(HEADERS, start=1):
        max_len = len(header)
        for row in sheet.iter_rows(min_row=2, min_col=index, max_col=index):
            value = row[0].value
            if value is not None:
                max_len = max(max_len, len(str(value)))
        sheet.column_dimensions[get_column_letter(index)].width = min(max_len + 2, 40)

    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer
